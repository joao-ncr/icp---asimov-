"""Importa leads.xlsx (CRM) como base de conhecimento: clientes, perdas e motivos.
A planilha tem colunas deslocadas nas linhas 'Perdida' (motivo cai em posicoes variaveis), entao a
classificacao e feita por TIPO de celula, nao por posicao."""
from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook
from app.core.utils import load_yaml, norm_name, jdump
from app.core.settings import get_settings

STAGES = {"Negociação", "Proposta Marcada", "Stand By", "Entrada", "Reunião Diagnóstica", "Primeiro Contato"}
SITUATIONS = {"Ganha", "Perdida"}
# ordem = prioridade quando houver mais de um motivo na linha
LOSS_CATEGORIES = [
    ("outside_icp", ["fora do icp"]),
    ("no_solvable_pain", ["nao possui dor", "não possui dor"]),
    ("disqualified", ["desqualificado"]),
    ("no_budget", ["nao possui orcamento", "não possui orçamento"]),
    ("contract_terms", ["termo do contrato"]),
    ("not_priority", ["nao e prioridade", "não é prioridade"]),
    ("no_response", ["sem resposta", "nao respondeu", "não respondeu", "nao consegui", "não consegui", "sem contato", "nao quis ouvir", "não quis ouvir"]),
]
MONTHS = {"janeiro": 1, "fevereiro": 2, "março": 3, "abril": 4, "maio": 5, "junho": 6, "julho": 7, "agosto": 8,
          "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12}
DATE_RE = re.compile(r"(\d{1,2}) de (\w+) de (\d{4})")


def categorize(reasons: list[str]) -> str | None:
    low = " | ".join(reasons).lower()
    for cat, keys in LOSS_CATEGORIES:
        if any(k in low for k in keys):
            return cat
    return "other" if reasons else None


def parse_rows(xlsx: Path) -> list[dict]:
    ann = load_yaml(get_settings().profile_dir / "leads_annotations.yaml")
    persons, overrides = set(ann["persons"]), ann["title_overrides"]
    segs = {k.lower(): v for k, v in ann["segments"].items()}
    ws = load_workbook(xlsx, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))[1:]
    out, n_person = [], 0
    for r in rows:
        cells = [c for c in r if c not in (None, "")]
        if len(cells) < 6:
            continue
        title = str(cells[0]).strip()
        situation = next((c for c in cells if c in SITUATIONS), None)
        if not situation:
            continue
        funnel = next((str(c) for c in cells if isinstance(c, str) and c.startswith("Funil")), None)
        value = next((float(c) for c in cells[3:5] if isinstance(c, (int, float))), 0.0)
        date = next((m for c in cells if isinstance(c, str) and (m := DATE_RE.search(c))), None)
        created = f"{date.group(3)}-{MONTHS.get(date.group(2), 0):02d}" if date else None
        reasons = [str(c) for c in cells[3:] if isinstance(c, str) and c not in STAGES and c not in SITUATIONS
                   and c != "Aberto" and not c.startswith("Funil") and not DATE_RE.search(c)]
        is_person = title in persons
        if is_person:
            n_person += 1
            name = f"Pessoa fisica #{n_person}"
        else:
            name = overrides.get(title, title)
        out.append({
            "name": name, "outcome": "won" if situation == "Ganha" else "lost", "value": value,
            "funnel": {"Funil Passivo": "passivo", "Funil Ativa": "ativo"}.get(funnel, funnel),
            "loss_reason": " | ".join(reasons) or None,
            "loss_category": categorize(reasons) if situation == "Perdida" else None,
            "segment": segs.get(name.lower(), "desconhecido"), "is_person": int(is_person), "created_month": created,
        })
    return out


def import_leads(conn, xlsx: Path) -> int:
    rows = parse_rows(xlsx)
    conn.execute("DELETE FROM known_lead")
    for r in rows:
        conn.execute("INSERT INTO known_lead(name,name_norm,domain,outcome,value,funnel,loss_reason,loss_category,segment,is_person,created_month)"
                     " VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                     (r["name"], norm_name(r["name"]), None, r["outcome"], r["value"], r["funnel"], r["loss_reason"],
                      r["loss_category"], r["segment"], r["is_person"], r["created_month"]))
    conn.commit()
    return len(rows)


def write_seed_csv(rows: list[dict], path: Path) -> None:
    import csv
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = ["name", "outcome", "value", "funnel", "loss_reason", "loss_category", "segment", "is_person", "created_month"]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({k: r[k] for k in cols})
